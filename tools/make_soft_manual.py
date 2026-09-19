# -*- coding: utf-8 -*-
"""生成软著登记用的《执课软件 V0.7.0 用户手册》.docx
页眉：执课软件 V0.7.0；页脚：页码；含封面/目录/正文 + 截图占位。
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SOFT_NAME = "执课软件"
VERSION = "V0.7.0"
HOLDER = "（著作权人姓名，登记时替换）"

DOC = Document()

# ---------- 全局样式 ----------
def set_run_font(run, east="宋体", size=11, bold=False, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if color:
        run.font.color.rgb = RGBColor(*color)

def para(text="", east="宋体", size=11, bold=False, align=None, space_after=6,
         space_before=0, indent=None, color=None):
    p = DOC.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.4
    if indent is not None:
        p.paragraph_format.first_line_indent = Cm(indent)
    if text:
        r = p.add_run(text)
        set_run_font(r, east=east, size=size, bold=bold, color=color)
    return p

def heading(text, level=1):
    p = DOC.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    if level == 1:
        set_run_font(r, east="黑体", size=15, bold=True)
    elif level == 2:
        set_run_font(r, east="黑体", size=12.5, bold=True)
    else:
        set_run_font(r, east="宋体", size=11.5, bold=True)
    return p

def shot(desc):
    """截图占位框"""
    p = DOC.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    # 占位：带边框的段落
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "dashed")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "AAAAAA")
        pbdr.append(e)
    pPr.append(pbdr)
    r = p.add_run(f"〔截图占位〕{desc}\n（请替换为软件实际运行界面截图）")
    set_run_font(r, east="宋体", size=10, color=(0x88, 0x88, 0x88))

# ---------- 页眉页脚 ----------
sec = DOC.sections[0]
sec.top_margin = Cm(2.4)
sec.bottom_margin = Cm(2.2)
sec.left_margin = Cm(2.5)
sec.right_margin = Cm(2.5)

hdr = sec.header
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hr = hp.add_run(f"{SOFT_NAME} {VERSION} 用户手册")
set_run_font(hr, east="宋体", size=9, color=(0x66, 0x66, 0x66))

ftr = sec.footer
fp = ftr.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("第 ")
set_run_font(r, east="宋体", size=9, color=(0x66, 0x66, 0x66))
# PAGE 域
run = fp.add_run()
fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
run._r.append(fld1); run._r.append(instr); run._r.append(fld2)
set_run_font(run, east="宋体", size=9, color=(0x66, 0x66, 0x66))
r2 = fp.add_run(" 页")
set_run_font(r2, east="宋体", size=9, color=(0x66, 0x66, 0x66))

# ---------- 封面 ----------
for _ in range(5):
    para("", space_after=0)
para(SOFT_NAME, east="黑体", size=30, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
para("（用户手册）", east="宋体", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=30)
para(f"软件版本：{VERSION}", east="宋体", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
para(f"著作权人：{HOLDER}", east="宋体", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
para("开发完成日期：2026 年 9 月", east="宋体", size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
DOC.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 目录 ----------
heading("目  录", 1)
toc = [
    "1  软件概述",
    "    1.1  软件简介",
    "    1.2  主要功能",
    "    1.3  运行环境",
    "2  安装与卸载",
    "    2.1  安装方法",
    "    2.2  卸载方法",
    "3  界面说明",
    "    3.1  主界面（课表页）",
    "    3.2  课表管理页",
    "    3.3  已添加课程页",
    "    3.4  设置页",
    "4  功能操作说明",
    "    4.1  切换周次",
    "    4.2  添加课程",
    "    4.3  编辑与删除课程",
    "    4.4  学期设置",
    "    4.5  每节课时间设置",
    "    4.6  课表背景设置",
    "    4.7  主题模式切换",
    "    4.8  从教务网站导入课表",
    "    4.9  导入课表文件",
    "    4.10  导出课表文件",
    "    4.11  桌面小组件",
    "5  常见问题",
    "6  隐私说明",
]
for t in toc:
    para(t, east="宋体", size=11, space_after=2)
DOC.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

# ---------- 1 概述 ----------
heading("1  软件概述", 1)
heading("1.1  软件简介", 2)
para("“执课”是一款运行于 Android 平台的课程表管理软件，致力于为用户提供干净、无广告的课程安排管理体验。"
     "软件具备零广告、零账号、零云同步的特点，所有课程数据均保存在用户设备本地，不采集、不上传任何个人信息。"
     "软件采用 Kotlin 语言与 Jetpack Compose 框架开发，界面简洁、操作直观，适合在校学生日常记录和查看课程安排。")
heading("1.2  主要功能", 2)
for t in [
    "（1）周课表展示：以周为单位展示周一至周日的课程安排，支持左右滑动与按钮切换周次；",
    "（2）课程管理：支持手动添加、编辑、删除课程，课程可按名称、时间、颜色进行区分；",
    "（3）学期与作息设置：支持自定义学期起始日期、总周数，以及每节课的上课时间；",
    "（4）教务导入：内置浏览器，可登录正方教务系统一键识别并导入课表；支持导入 HTML/MHTML 课表文件与备份文件；",
    "（5）数据备份：支持将课表数据导出为 JSON 备份文件，便于迁移与保存；",
    "（6）主题切换：提供浅色、深色、跟随系统三种主题模式；",
    "（7）桌面小组件：提供周视图、日视图、近日课程、今日课程四种桌面小组件，支持上下滑动查看全部课程。",
]:
    para(t, indent=0.74)
heading("1.3  运行环境", 2)
para("操作系统：Android 7.0（API 24）及以上版本；", indent=0.74)
para("存储空间：约 30 MB；", indent=0.74)
para("网络环境：仅在“从教务网站导入课表”时需访问互联网，其余功能均可离线使用。", indent=0.74)

# ---------- 2 安装卸载 ----------
heading("2  安装与卸载", 1)
heading("2.1  安装方法", 2)
para("将软件安装包（APK 文件）复制到 Android 设备后点击打开，系统会提示安装，"
     "在系统权限确认界面点击“安装”即可完成。首次安装时系统可能提示“允许安装未知来源应用”，"
     "按提示在系统设置中授权后即可继续安装。")
shot("安装包安装确认界面")
heading("2.2  卸载方法", 2)
para("在系统“设置 - 应用管理”中找到“执课”应用，点击进入详情页后选择“卸载”，"
     "确认后即可卸载。卸载会同时清除本机保存的课程数据，请提前导出备份。")

# ---------- 3 界面说明 ----------
heading("3  界面说明", 1)
heading("3.1  主界面（课表页）", 2)
para("主界面为周课表视图：顶部显示当前学期名称、周次与日期范围，下方为周一至周日七列的课程网格。"
     "右上角提供“管理”与“设置”入口，右下角为“添加课程”按钮。")
shot("主界面 - 周课表视图")
heading("3.2  课表管理页", 2)
para("课表管理页用于管理多个课表（学期），支持新建课表、切换当前课表、重命名与删除课表。")
shot("课表管理页")
heading("3.3  已添加课程页", 2)
para("已添加课程页按课程名称聚合展示全部课程，支持按课程批量删除、统一改名、改色与修改时间。")
shot("已添加课程页")
heading("3.4  设置页", 2)
para("设置页按“这张课表”“课程数据”“外观”“关于”分组，集中管理学期、作息、背景、导入导出、主题等设置项。")
shot("设置页")

# ---------- 4 功能 ----------
heading("4  功能操作说明", 1)
heading("4.1  切换周次", 2)
para("在课表主界面，通过以下任一方式切换周次："
     "（1）在课表区域左右滑动；（2）点击顶部周次两侧的左右箭头；（3）点击顶部周次区域打开“跳转周次”面板，"
     "可网格点选、拖动滑条或直接输入周次数字。")
shot("跳转周次面板")
heading("4.2  添加课程", 2)
para("点击主界面右下角“添加课程”按钮，在弹出的编辑窗口中填写课程名称、教师、上课地点、"
     "上课周次（星期）、起始节次与结束节次等信息，点击“保存”即可将课程添加到对应课表。")
shot("添加课程编辑窗口")
heading("4.3  编辑与删除课程", 2)
para("点击课表中已有的课程卡片，可打开编辑窗口修改课程信息；在编辑窗口点击“删除”即可删除该课程。")
shot("课程编辑窗口")
heading("4.4  学期设置", 2)
para("在设置页点击“学期设置”，可设置第一周的起始日期与学期总周数。起始日期决定周次计算基准，"
     "总周数决定学期课表的时间跨度。")
shot("学期设置窗口")
heading("4.5  每节课时间设置", 2)
para("在设置页点击“每节课时间”，进入作息时间编辑页。可设置每一节课的起止时间，支持按“时/分”分别填写、"
     "在中间插入新的一节、以及开启“只填写开始时间”模式。")
shot("每节课时间编辑页")
heading("4.6  课表背景设置", 2)
para("在设置页点击“课表背景”，可从相册选择一张图片作为课表背景，并调节蒙层透明度；"
     "点击“清除背景”可恢复默认背景。")
shot("课表背景设置窗口")
heading("4.7  主题模式切换", 2)
para("在设置页“外观”分组下，可在“浅色 / 深色 / 跟随系统”三种主题模式间切换，切换立即生效并同步到全部页面。")
shot("主题模式切换")
heading("4.8  从教务网站导入课表", 2)
para("在设置页点击“从教务网站导入”，打开内置浏览器并跳转至教务系统登录页。用户自行登录后，"
     "软件会自动识别页面中的课程表并高亮提示，点击“一键导入”即可将课表解析并导入，导入前会弹出预览确认。")
shot("内置浏览器导入教务课表")
heading("4.9  导入课表文件", 2)
para("在设置页点击“导入课表文件”，可从文件管理器选择 HTML / MHTML 课表文件或此前导出的 JSON 备份文件，"
     "软件解析后弹出预览，确认后导入。也可通过其他应用“分享”或文件管理器“打开方式”将文件发送给本软件导入。")
shot("导入课表文件")
heading("4.10  导出课表文件", 2)
para("在设置页点击“导出课表文件”，可将当前课表数据导出为 JSON 备份文件，用于备份或在其他设备上导入。")
shot("导出课表文件")
heading("4.11  桌面小组件", 2)
para("在系统桌面长按空白处进入“小组件”列表，选择“执课”的四种小组件之一添加到桌面："
     "周视图（4×2）、日视图（4×2）、近日课程（4×2）、今日课程（2×2）。"
     "当课程数量超过一屏时，可上下滑动查看其余课程。")
shot("桌面小组件")

# ---------- 5 FAQ ----------
heading("5  常见问题", 1)
faq = [
    ("Q1：为什么添加课程时“节次”只能输入两位数？",
     "A：旧版本曾存在输入框不允许清空的问题，本版本已修复。填写节次时可直接清空后输入任意正整数，"
     "退出输入框后自动按有效范围修正。"),
    ("Q2：导入教务课表失败怎么办？",
     "A：请确认已正确登录教务系统并停留在课表页面；如仍失败，可先将课表页面“另存为 HTML 文件”，"
     "再通过“导入课表文件”功能导入。"),
    ("Q3：课程数据会同步到云端吗？",
     "A：不会。本软件零账号、零云同步，全部数据仅保存在本机。建议定期使用“导出课表文件”进行备份。"),
    ("Q4：桌面小组件不显示课程怎么办？",
     "A：请在软件内打开一次主界面以触发刷新；若仍不显示，可移除小组件后重新添加。"),
]
for q, a in faq:
    para(q, bold=True, space_after=2)
    para(a, indent=0.74, space_after=6)

# ---------- 6 隐私 ----------
heading("6  隐私说明", 1)
para("本软件坚持最小化权限与本地化存储原则：除“从教务网站导入课表”时需要访问互联网外，"
     "不申请任何多余系统权限；不采集、不存储、不上传用户的任何个人身份信息、账号密码与课程数据；"
     "在导入教务课表过程中，登录密码由教务网站自行处理，不经过本软件，关闭导入页面即清除访问记录。")

# ---------- 保存 ----------
out_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs", "软著")
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, f"{SOFT_NAME}{VERSION}用户手册.docx")
DOC.save(out_path)
print("SAVED:", out_path)
